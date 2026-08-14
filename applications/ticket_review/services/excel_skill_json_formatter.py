from __future__ import annotations

import asyncio
import json
import logging
import re
from dataclasses import replace
from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path
from typing import Any

from claude_agent_sdk import ClaudeAgentOptions, ClaudeSDKClient, ResultMessage
from openpyxl import load_workbook
from pydantic import BaseModel, Field

try:
    from json_repair import repair_json
except ModuleNotFoundError:  # Optional fallback; SDK structured output is the primary path.
    repair_json = None

from configure import PROJECT_CONFIG

from applications.ticket_review.schemas.push_ticket import (
    is_review_record,
    is_supported_review_table,
)

from applications.ticket_review.services.department_tool import (
    DEPARTMENT_ALLOWED_TOOL,
    DEPARTMENT_MCP_SERVER_KEY,
    create_department_mcp_server,
)
from applications.ticket_review.services.ping_tool import (
    PING_ALLOWED_TOOL,
    PING_MCP_SERVER_KEY,
    create_ping_mcp_server,
)

logger = logging.getLogger(__name__)

PROJECT_ROOT = Path(PROJECT_CONFIG.WORKSPACE_DIR)
COCO_SKILL_DIR = PROJECT_ROOT / ".claude" / "skills" / "coco"

FIELD_ALIASES = {
    "ticket_id": ["id", "ticket_id", "ticketid", "工单id", "申请单id", "记录id"],
    "instance_id": ["instance_id", "instanceid", "实例id", "流程实例id"],
    "ticket_no": ["extra1", "ticket_no", "ticketno", "工单号", "工单编号", "申请单号", "单号"],
    "work_type": ["work_type", "worktype", "工单类型", "申请类型", "流程类型", "事件类型"],
    "title": ["title", "subject", "工单标题", "申请标题", "标题", "主题"],
    "app_system": ["app_system", "appsystem", "应用系统", "系统名称", "所属系统", "paso系统"],
    "user_org": ["user_org", "userorg", "用户组织", "申请人部门", "所属部门", "paso所属部门", "组织名称", "部门名称"],
    "content": ["content", "description", "工单内容", "申请内容", "工单描述", "需求描述", "详情", "描述", "申请原因"],
}

# JSON push fields share the same normalized review input as Excel aliases.
FIELD_ALIASES["content"].extend(["work_content", "malfunction_appearance"])


class CocoJsonMetadata(BaseModel):
    source_file: str = Field(description="Excel source file name")
    sheet_count: int = Field(description="Number of sheets parsed from the source file")
    row_count: int = Field(description="Number of data rows parsed from the source file")
    skill_name: str = Field(description="Skill used to generate the response")
    classification_docs: list[str] = Field(description="Classification documents used")
class CocoTicketReviewResult(BaseModel):
    row_number: int = Field(description="Source Excel row number")
    ticket_id: str = Field(description="Ticket ID from Excel id field")
    instance_id: str = Field(description="Ticket instance ID from Excel instance_id field")
    ticket_no: str = Field(description="Ticket number, usually extra1")
    ticket_type: str = Field(description="工单类型：服务请求类或故障类")
    ticket_center: str = Field(description="工单所属中心：上海、苏州或无法判断")
    title: str = Field(description="Ticket title")
    work_type: str = Field(description="Original work_type from Excel")
    app_system: str = Field(description="Application system from Excel")
    matched_keyword: str = Field(description="Matched keyword in classification document")
    primary_category: str = Field(description="一级分类 or 分类 from classification document")
    secondary_category: str = Field(description="二级分类 when available, empty string otherwise")
    screening_condition: str = Field(description="Matched 初筛判定条件")
    required_fields: list[str] = Field(description="Required fields from 涉及字段信息（必填）")
    missing_fields: list[str] = Field(description="Required fields missing from ticket content")
    information_complete: bool = Field(description="Whether all required fields are present")
    operation_party: str = Field(description="工单操作方：系统科、对应中心环境组、退回申请人、人工处理")
    assignee: str = Field(description="Final assignee. Must exactly follow coco skill dispatch logic")
    review_result: str = Field(description="预审结论：通过、退回申请人、人工处理")
    reason: str = Field(description="Concise Chinese reasoning based on the skill and classification document")
    source_sheet: str = Field(default="", description="Original Excel sheet name")
    raw_data: dict[str, Any] = Field(default_factory=dict, description="Unmodified original Excel row")


class CocoJsonResponse(BaseModel):
    metadata: CocoJsonMetadata
    tickets: list[CocoTicketReviewResult]


class ExcelSkillJsonFormatter:
    """Parse Excel rows and ask Claude to format them as coco skill JSON."""

    async def stream_excel_batches(self, excel_bytes: bytes, filename: str, batch_size: int = 5):
        """Yield one LLM result at a time while keeping batches strictly bounded."""
        workbook_data = self._parse_excel(excel_bytes, filename)
        batches = list(self._split_workbook_data(workbook_data, batch_size))
        metadata: dict[str, Any] = {
            "source_file": workbook_data["file_name"],
            "sheet_count": workbook_data["sheet_count"],
            "row_count": workbook_data["row_count"],
            "skill_name": "coco",
            "classification_docs": [],
        }
        yield {
            "type": "start",
            "metadata": metadata,
            "totalRows": workbook_data["row_count"],
            "totalBatches": len(batches),
            "batchSize": batch_size,
        }

        classification_docs: list[str] = []
        for index, batch_data in enumerate(batches, start=1):
            batch_json = await self._format_workbook_data_once(batch_data)
            batch_result = json.loads(batch_json)
            tickets = batch_result.get("tickets", [])
            self._attach_original_rows(tickets, batch_data)
            for doc in batch_result.get("metadata", {}).get("classification_docs", []):
                if doc not in classification_docs:
                    classification_docs.append(doc)
            metadata["classification_docs"] = classification_docs
            yield {
                "type": "batch",
                "batchIndex": index,
                "totalBatches": len(batches),
                "processedRows": min(index * batch_size, workbook_data["row_count"]),
                "totalRows": workbook_data["row_count"],
                "tickets": tickets,
                "metadata": metadata,
            }
    async def stream_push_payload(
        self,
        payload: dict[str, Any],
        batch_size: int = 5,
        start_after: int = 0,
    ):
        """Review FW038/FW039 JSON records through the existing LLM pipeline."""
        workbook_data = self._push_payload_to_workbook_data(payload)
        batches = list(self._split_workbook_data(workbook_data, batch_size))
        metadata: dict[str, Any] = {
            "source_file": workbook_data["file_name"],
            "sheet_count": workbook_data["sheet_count"],
            "row_count": workbook_data["row_count"],
            "skill_name": "coco",
            "classification_docs": [],
        }
        yield {
            "type": "start",
            "metadata": metadata,
            "totalRows": workbook_data["row_count"],
            "totalBatches": len(batches),
            "batchSize": batch_size,
        }

        classification_docs: list[str] = []
        for index, batch_data in enumerate(batches, start=1):
            processed_rows = min(index * batch_size, workbook_data["row_count"])
            if processed_rows <= start_after:
                continue
            batch_json = await self._format_workbook_data_once(batch_data)
            batch_result = json.loads(batch_json)
            tickets = batch_result.get("tickets", [])
            self._attach_original_rows(tickets, batch_data)
            for doc in batch_result.get("metadata", {}).get("classification_docs", []):
                if doc not in classification_docs:
                    classification_docs.append(doc)
            metadata["classification_docs"] = classification_docs
            yield {
                "type": "batch",
                "batchIndex": index,
                "totalBatches": len(batches),
                "processedRows": processed_rows,
                "totalRows": workbook_data["row_count"],
                "tickets": tickets,
                "metadata": metadata,
            }

    def _push_payload_to_workbook_data(self, payload: dict[str, Any]) -> dict[str, Any]:
        """Represent JSON items as virtual sheets without dropping original fields."""
        sheets: list[dict[str, Any]] = []
        row_number = 0
        # for item in payload.get("items", []):
        #     if not isinstance(item, dict):
        #         continue
        #     table_code = str(item.get("tableCode") or "").upper()
        #     records = item.get("records")
        #     if table_code not in {"FW038", "FW039"} or not isinstance(records, list) or not records:
        #         continue

        #     table_name = str(item.get("tableName") or table_code)
        #     rows: list[dict[str, Any]] = []
        #     headers: list[str] = []
        #     for source in records:
        #         if not isinstance(source, dict):
        #             continue
        #         row_number += 1
        #         record = dict(source)
        for item in payload.get("items", []):
            if not isinstance(item, dict):
                continue

            table_code = str(item.get("tableCode") or "").strip().upper()
            records = item.get("records")

            if (
                not is_supported_review_table(table_code)
                or not isinstance(records, list)
                or not records
            ):
                continue

            table_name = str(item.get("tableName") or table_code)
            rows: list[dict[str, Any]] = []
            headers: list[str] = []

            for source in records:
                if not is_review_record(table_code, source):
                    continue

                row_number += 1
                record = dict(source)
                record["work_type"] = "故障上报" if table_code == "FW038" else "服务请求"
                for key in record:
                    if key not in headers:
                        headers.append(key)
                record["_table_code"] = table_code
                record["_table_name"] = table_name
                record["_row_number"] = row_number
                record["_normalized"] = self._normalize_record(record)
                rows.append(record)

            if rows:
                sheets.append(
                    {
                        "sheet_name": f"{table_code}-{table_name}",
                        "headers": headers,
                        "rows": rows,
                    }
                )

        if not sheets:
            raise ValueError("No non-empty FW038/FW039 records for 开发测试审批组 were found")
        message_id = str(payload.get("messageId") or "push-task")
        return {
            "file_name": f"{message_id}.json",
            "sheet_count": len(sheets),
            "row_count": row_number,
            "sheets": sheets,
        }
    @staticmethod
    def _attach_original_rows(tickets: list[dict[str, Any]], batch_data: dict[str, Any]) -> None:
        """Attach source rows deterministically instead of asking the LLM to reproduce them."""
        source_rows = [
            (sheet["sheet_name"], row)
            for sheet in batch_data["sheets"]
            for row in sheet["rows"]
        ]
        unused = list(source_rows)
        for ticket in tickets:
            row_number = ticket.get("row_number")
            match_index = next(
                (index for index, (_, row) in enumerate(unused) if row.get("_row_number") == row_number),
                0 if unused else None,
            )
            if match_index is None:
                continue
            # sheet_name, source_row = unused.pop(match_index)
            # ticket["source_sheet"] = sheet_name
            # ticket["raw_data"] = {
            #     key: value for key, value in source_row.items() if not key.startswith("_")
            # }
            sheet_name, source_row = unused.pop(match_index)

            # table_code由后端根据原始报文确定，不让LLM判断。
            ticket["table_code"] = str(
                source_row.get("_table_code") or ""
            ).strip().upper()

            ticket["source_sheet"] = sheet_name
            ticket["raw_data"] = {
                key: value
                for key, value in source_row.items()
                if not key.startswith("_")
            }
            normalized = source_row.get("_normalized", {})
            for field in ("ticket_id", "instance_id", "ticket_no", "title", "work_type", "app_system"):
                if not ticket.get(field) and normalized.get(field) not in (None, ""):
                    ticket[field] = str(normalized[field])

    async def format_excel(self, excel_bytes: bytes, filename: str) -> str:
        logger.info("[1/4] 开始解析 Excel: %s", filename)
        workbook_data = self._parse_excel(excel_bytes, filename)
        batch_size = PROJECT_CONFIG.TICKET_COCO_BATCH_SIZE
        if batch_size > 0 and workbook_data["row_count"] > batch_size:
            return await self._format_excel_in_batches(workbook_data, batch_size)
        return await self._format_workbook_data_once(workbook_data)

    async def _format_excel_in_batches(self, workbook_data: dict[str, Any], batch_size: int) -> str:
        logger.info(
            "[2/4] Excel 解析完成: %s 行, %s 个 sheet，启用批处理，每批 %s 行",
            workbook_data["row_count"],
            workbook_data["sheet_count"],
            batch_size,
        )
        combined: dict[str, Any] = {
            "metadata": {
                "source_file": workbook_data["file_name"],
                "sheet_count": workbook_data["sheet_count"],
                "row_count": workbook_data["row_count"],
                "skill_name": "coco",
                "classification_docs": [],
            },
            "tickets": [],
        }
        classification_docs: list[str] = []

        batches = list(self._split_workbook_data(workbook_data, batch_size))
        for index, batch_data in enumerate(batches, start=1):
            logger.info(
                "[批处理] 开始处理第 %s/%s 批，%s 行",
                index,
                len(batches),
                batch_data["row_count"],
            )
            batch_json = await self._format_workbook_data_once(batch_data)
            batch_result = json.loads(batch_json)
            combined["tickets"].extend(batch_result.get("tickets", []))
            for doc in batch_result.get("metadata", {}).get("classification_docs", []):
                if doc not in classification_docs:
                    classification_docs.append(doc)

        combined["metadata"]["classification_docs"] = classification_docs or [
            "服务请求类工单场景分类.md",
            "故障类工单场景分类.md",
        ]
        combined = self._apply_coco_business_corrections(combined)
        validated = CocoJsonResponse.model_validate(combined)
        logger.info("[4/4] 批处理完成，共合并 %s 条结果", len(validated.tickets))
        return validated.model_dump_json(indent=2)
    def _split_workbook_data(self, workbook_data: dict[str, Any], batch_size: int):
        rows_with_sheet: list[tuple[dict[str, Any], dict[str, Any]]] = []
        for sheet in workbook_data["sheets"]:
            for row in sheet["rows"]:
                rows_with_sheet.append((sheet, row))

        for start in range(0, len(rows_with_sheet), batch_size):
            batch_items = rows_with_sheet[start : start + batch_size]
            sheet_map: dict[str, dict[str, Any]] = {}
            for sheet, row in batch_items:
                key = sheet["sheet_name"]
                if key not in sheet_map:
                    sheet_map[key] = {
                        "sheet_name": sheet["sheet_name"],
                        "headers": sheet["headers"],
                        "rows": [],
                    }
                sheet_map[key]["rows"].append(row)

            yield {
                "file_name": workbook_data["file_name"],
                "sheet_count": len(sheet_map),
                "row_count": len(batch_items),
                "sheets": list(sheet_map.values()),
            }

    async def _format_workbook_data_once(self, workbook_data: dict[str, Any]) -> str:
        prompt = self._build_prompt(workbook_data)
        logger.info(
            "[2/4] Excel 解析完成: %s 行, %s 个 sheet, prompt 大小 %s bytes",
            workbook_data["row_count"],
            workbook_data["sheet_count"],
            len(prompt.encode("utf-8")),
        )

        schema_text = json.dumps(CocoJsonResponse.model_json_schema(), ensure_ascii=False)
        base_options = ClaudeAgentOptions(
            max_turns=80,
            system_prompt=(
                "你是测试环境工单预审人员。必须先使用 Skill 工具加载项目级 coco skill，"
                "并使用 Read 工具读取该 skill 指定的分类规则文档；禁止跳过 skill 自行判断。"
                "分类规则读取顺序：1) 先读取精确分类文档（服务请求类→服务请求类工单场景分类.md，故障类→故障类工单场景分类.md）；"
                "2) 能唯一命中精确规则时禁止使用兜底；"
                "3) 无法唯一命中精确规则但对象明确时，再读取对应兜底文档（服务请求类→服务请求类工单兜底分类.md，故障类→故障类工单兜底分类.md）；"
                "4) 对象不明确时输出人工处理。"
                "历史Excel已被离线总结进别名和对象映射，线上不直接读取历史Excel。"
                "禁止仅凭操作方相同直接通过，禁止用必填字段反推对象或场景，禁止为降低人工处理率猜测对象或编造分类。"
                "本接口是后台自动化任务，不要询问用户，不要写文件，只返回 JSON 字符串。"
                "你拥有resolve_center_by_org_name工具。必须先完成工单分类；只有一级分类为堡垒机、软负载或DFTM，"
                "并且user_org不为空且不包含上海或苏州时，才必须调用该工具。其他情况禁止调用。"
                "调用时必须把Excel原始user_org原样传给工具，以工具返回的center作为ticket_center；"
                "工具返回无法判断时不得猜测上海或苏州。禁止编写或执行任何SQL；带有 ***分行的，调用工具之后没法判断中心的话归属为 上海"
                "数据库查询只能通过resolve_center_by_org_name工具完成。"
                "你还拥有check_ip_reachable工具，用于检测IP的ICMP连通性。"
                "处理FW038故障上报工单时，如果某条工单的malfunction_appearance、title或其他原始字段中包含明确、完整的IP地址，"
                "必须调用check_ip_reachable工具检测该IP。调用参数格式为{\"ip\": \"原始IP\"}。"
                "不得猜测或编造IP，同一条工单中的同一个IP只调用一次。将工具返回的检测结果写入reason。"
                "reachable=false只表示未收到ICMP响应，可能是网络不可达、服务器未启动或防火墙禁止ICMP，"
                "不能仅凭ping失败直接判断服务器宕机、退回申请人或转人工处理。FW039不调用该工具。"
                f"必须严格按以下 JSON Schema 返回：{schema_text}"
            ),
            mcp_servers={
                DEPARTMENT_MCP_SERVER_KEY: create_department_mcp_server(),
                PING_MCP_SERVER_KEY: create_ping_mcp_server(),
            },
            allowed_tools=["Skill", "Read", DEPARTMENT_ALLOWED_TOOL, PING_ALLOWED_TOOL],
            # CC Switch writes the active provider into ~/.claude/settings.json.
            # Project settings are required for discovery of .claude/skills/coco.
            setting_sources=["user", "project"],
            disallowed_tools=[
                "AskUserQuestion", "Glob", "Grep", "Write", "Edit",
                "Bash", "WebFetch", "WebSearch", "NotebookEdit", "Task",
            ],
            cwd=str(PROJECT_ROOT),
            output_format={
                "type": "json_schema",
                "schema": CocoJsonResponse.model_json_schema(),
            },
        )
        sdk_pool = PROJECT_CONFIG.TICKET_SDK_MODEL_POOL.strip()
        models = sdk_pool.split(";") if sdk_pool else PROJECT_CONFIG.TICKET_MODEL_POOL.split(";")
        default_model = PROJECT_CONFIG.TICKET_ANTHROPIC_MODEL.strip()
        if default_model and default_model not in models:
            models.append(default_model)
        logger.info("[3/4] SDK 模型池: %s", models)

        last_error = ""
        retry_count = max(1, PROJECT_CONFIG.TICKET_MODEL_RETRY_COUNT)
        for index, model in enumerate(models, start=1):
            if not model:
                continue
            for attempt in range(1, retry_count + 1):
                logger.info(
                    "[3/4] SDK 尝试模型 %s/%s: %s，第 %s/%s 次",
                    index,
                    len(models),
                    model,
                    attempt,
                    retry_count,
                )
                options = replace(base_options, model=model)
                try:
                    result_text = await asyncio.wait_for(
                        self._run_session(options, prompt),
                        timeout=PROJECT_CONFIG.TICKET_MODEL_TIMEOUT,
                    )
                    logger.info("[4/4] SDK 模型 %s 调用成功", model)
                    return self._normalize_json_string(result_text)
                except asyncio.TimeoutError:
                    last_error = f"模型 {model} 执行超时（{PROJECT_CONFIG.TICKET_MODEL_TIMEOUT}秒）"
                    logger.warning(last_error)
                except Exception as exc:
                    last_error = f"模型 {model} 执行失败: {exc}"
                    logger.warning(last_error)
                if attempt < retry_count:
                    await asyncio.sleep(max(0, PROJECT_CONFIG.TICKET_MODEL_RETRY_DELAY))

        logger.error("[4/4] 所有调用方式均失败")
        raise RuntimeError(f"Claude Agent SDK所有模型均失败，最后错误：{last_error}")

    def _parse_excel(self, excel_bytes: bytes, filename: str) -> dict[str, Any]:
        workbook = load_workbook(BytesIO(excel_bytes), read_only=True, data_only=True)
        sheets: list[dict[str, Any]] = []
        total_rows = 0

        for worksheet in workbook.worksheets:
            rows = list(worksheet.iter_rows(values_only=True))
            if not rows:
                continue

            header_index = self._detect_header_row_index(rows)
            if header_index is None:
                continue

            raw_headers = [
                self._header_name(value, index)
                for index, value in enumerate(rows[header_index], start=1)
            ]
            headers = self._make_headers_unique(raw_headers)

            records: list[dict[str, Any]] = []
            for row_number, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
                values = [self._cell_value(value) for value in row]
                if not any(value not in ("", None) for value in values):
                    continue

                record = {
                    headers[index]: values[index] if index < len(values) else ""
                    for index in range(len(headers))
                }
                record["_normalized"] = self._normalize_record(record)
                record["_row_number"] = row_number
                records.append(record)

            if records:
                total_rows += len(records)
                sheets.append(
                    {
                        "sheet_name": worksheet.title,
                        "headers": headers,
                        "rows": records,
                    }
                )

        if not sheets:
            raise ValueError("Excel 文件中未解析到有效数据行")

        return {
            "file_name": filename,
            "sheet_count": len(sheets),
            "row_count": total_rows,
            "sheets": sheets,
        }
    def _first_non_empty_row_index(self, rows: list[tuple[Any, ...]]) -> int | None:
        for index, row in enumerate(rows):
            if any(value not in ("", None) for value in row):
                return index
        return None

    def _detect_header_row_index(self, rows: list[tuple[Any, ...]]) -> int | None:
        """Find the real header even when a report title precedes it."""
        candidates: list[tuple[int, int]] = []
        known_aliases = {
            self._canonical_header(alias)
            for aliases in FIELD_ALIASES.values()
            for alias in aliases
        }
        for index, row in enumerate(rows[:20]):
            values = [str(value).strip() for value in row if value not in (None, "")]
            if not values:
                continue
            alias_hits = sum(
                1
                for value in values
                if any(
                    alias == self._canonical_header(value)
                    or (len(alias) >= 3 and alias in self._canonical_header(value))
                    for alias in known_aliases
                )
            )
            candidates.append((alias_hits * 100 + len(values), index))
        return max(candidates)[1] if candidates else None

    @staticmethod
    def _canonical_header(value: Any) -> str:
        return re.sub(r"[\s_\-—:：/\\（）()\[\]【】]+", "", str(value or "").strip().lower())

    @staticmethod
    def _make_headers_unique(headers: list[str]) -> list[str]:
        counts: dict[str, int] = {}
        result: list[str] = []
        for header in headers:
            counts[header] = counts.get(header, 0) + 1
            result.append(header if counts[header] == 1 else f"{header}__{counts[header]}")
        return result

    def _normalize_record(self, record: dict[str, Any]) -> dict[str, Any]:
        canonical_items = [
            (self._canonical_header(header), value)
            for header, value in record.items()
            if not header.startswith("_")
        ]
        normalized: dict[str, Any] = {}
        for target, aliases in FIELD_ALIASES.items():
            canonical_aliases = [self._canonical_header(alias) for alias in aliases]
            value = next(
                (value for header, value in canonical_items if header in canonical_aliases and value not in (None, "")),
                None,
            )
            if value is None:
                value = next(
                    (
                        value
                        for header, value in canonical_items
                        if value not in (None, "")
                        and any(len(alias) >= 3 and alias in header for alias in canonical_aliases)
                    ),
                    "",
                )
            normalized[target] = value
        return normalized

    def _header_name(self, value: Any, index: int) -> str:
        normalized = str(value).strip() if value is not None else ""
        return normalized or f"column_{index}"

    def _cell_value(self, value: Any) -> Any:
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d %H:%M:%S")
        if isinstance(value, date):
            return value.strftime("%Y-%m-%d")
        if isinstance(value, time):
            return value.strftime("%H:%M:%S")
        return value
    def _build_prompt(self, workbook_data: dict[str, Any]) -> str:
        workbook_json = json.dumps(workbook_data, ensure_ascii=False, indent=2)
        return f"""
字段兼容说明：真实 Excel 的表头可能与样例不同。每条 record 中：
- 原始表头和原始值全部保留，不得因为字段较多而忽略工单内容。
- `_normalized` 是后端根据表头别名生成的标准字段视图，优先使用其中的
  ticket_id、instance_id、ticket_no、work_type、title、app_system、user_org、content。
- `_normalized` 为空的标准字段，应继续结合该行其他原始字段判断，不得编造。
- 必填信息检查必须搜索该行全部原始字段的名称和值，不能只检查模拟 Excel 中出现过的列。
中心归属工具规则：必须先分类。仅当一级分类是”堡垒机”、”软负载”或”DFTM”，并且
`_normalized.user_org`非空且不包含”上海/苏州”时，必须调用 resolve_center_by_org_name，
输入完整原始user_org。分类不属于这三类、user_org为空、或user_org已包含上海/苏州时禁止调用。
工具结果为上海或苏州时必须原样写入 ticket_center；结果为无法判断时不得自行猜测。
IP 连通性检测规则：处理FW038故障上报工单时，如果某条工单的 malfunction_appearance、
title 或其他原始字段中包含明确、完整的 IP 地址，必须调用 check_ip_reachable 工具检测该IP。
调用参数格式为 {{"ip": "原始IP"}}。不得猜测或编造IP，同一条工单中的同一个IP只调用一次。
将工具返回的检测结果写入 reason。reachable=false 只表示未收到ICMP响应，可能是网络不可达、
服务器未启动或防火墙禁止ICMP，不能仅凭 ping 失败直接判断服务器宕机、退回申请人或转人工处理。
FW039 不调用该工具。
请根据下面 Excel 解析结果，结合 coco skill 的要求，输出格式化 JSON 字符串。

要求：
1. 必须先依据 work_type 判断工单类型：故障上报按”故障类”处理，读取`故障类工单场景分类.md`和`故障类工单兜底分类.md`；服务请求按”服务请求类”处理，读取`服务请求类工单场景分类.md`和`服务请求类工单兜底分类.md`。
2. 分类匹配优先级：
   a) 优先匹配精确分类规则中的”关键字/分类/初筛判定条件”；
   b) 能唯一命中精确规则时，必须使用精确规则的分类、必填字段和操作方，禁止使用兜底；
   c) 无法唯一命中精确规则，但对象已明确时，读取对应兜底文档，选择该对象的兜底规则；
   d) 对象不明确、候选规则冲突或存在无法拆分的多操作方诉求时，输出”人工处理”。
3. 对象确定依据：标题、content、work_content、malfunction_appearance及所有原始字段中出现标准对象名或已确认别名。仅有IP/端口/连接失败等通用信息不能确定具体对象。端口只能作为辅助证据（如3306可辅助MySQL判断，不能单独确定数据库）。
4. 历史Excel已被离线总结进别名映射和兜底规则，线上不直接读取历史Excel文件。
5. 必须从工单内容判断工单所属中心。若 user_org、title 或工单内容包含”上海”，所属中心为上海；包含”苏州”，所属中心为苏州；无法判断则填”无法判断”。
6. 必填字段检查：必须搜索该行全部原始字段的名称和值。”至少满足其中一个”的替代关系要明确识别。字段完整则继续派单，字段缺失则退回申请人。
7. 禁止为必填字段编造原文不存在的信息。”使用默认端口”不能擅自转换为具体端口，除非规则明确允许。
8. 操作方映射必须严格遵循 coco skill：
   - 系统科 -> assignee 输出”测试管理审批组”
   - 对应中心环境组 + 上海 + 堡垒机类 -> “张嘉瑾”
   - 对应中心环境组 + 上海 + DFTM类 -> “左义坤”
   - 对应中心环境组 + 苏州 + 堡垒机类 -> “孙坤铭”
   - 对应中心环境组 + 苏州 + DFTM类 -> “余守业”
   - 退回申请人 -> “退回申请人”
   - 人工处理 -> “人工处理”
9. JSON 顶层必须是对象，字段固定为 metadata、tickets。
10. 每一行 Excel 数据都必须在 tickets 中输出一条预审和分派结果。
11. metadata.classification_docs 必须包含实际使用的分类文档文件名（精确规则和兜底规则文档名均需记录）。
12. 严格按照 ClaudeAgentOptions.output_format 中的 JSON Schema 返回结构化 JSON。
Excel 解析结果：
{workbook_json}
"""

    async def _run_session(self, options: ClaudeAgentOptions, prompt: str) -> str:
        logger.info("[SDK] 连接模型: %s", options.model)
        client = ClaudeSDKClient(options=options)
        try:
            await client.connect(prompt=prompt)
            logger.info("[SDK] 连接成功，等待响应...")
            result_text = ""
            structured_output: Any = None
            message_count = 0
            async for message in client.receive_response():
                message_count += 1
                if isinstance(message, ResultMessage):
                    logger.info(
                        "[SDK] 收到结果消息, is_error=%s, api_status=%s, result长度=%s",
                        message.is_error,
                        message.api_error_status,
                        len(message.result or ""),
                    )
                    if message.structured_output is not None:
                        # Some compatible proxy models emit a valid StructuredOutput
                        # and then fail on a redundant final turn. The validated
                        # business result is still usable and must not be discarded.
                        structured_output = message.structured_output
                    if message.is_error and structured_output is None:
                        error_detail = "; ".join(message.errors) if message.errors else (
                            f"subtype={message.subtype or 'unknown'}, "
                            f"stop_reason={message.stop_reason or 'unknown'}, "
                            f"session_id={message.session_id or 'unknown'}"
                        )
                        status_hint = f" (HTTP {message.api_error_status})" if message.api_error_status else ""
                        raise RuntimeError(
                            f"Claude API 调用失败{status_hint}: {error_detail}"
                        )
                    if message.is_error:
                        logger.warning(
                            "[SDK] 模型结束状态异常，但已收到结构化结果，继续校验并保存: session_id=%s",
                            message.session_id,
                        )
                    result_text = message.result or ""
            logger.info("[SDK] 响应流结束, 共收到 %s 条消息", message_count)
            if structured_output is not None:
                parsed_output = CocoJsonResponse.model_validate(structured_output)
                return parsed_output.model_dump_json(indent=2)
            if not result_text:
                raise RuntimeError("Claude Agent SDK returned empty result")
            return result_text
        except Exception:
            logger.exception("[SDK] 调用失败")
            raise
        finally:
            await client.disconnect()
    def _normalize_json_string(self, text: str) -> str:
        if not text or not text.strip():
            raise RuntimeError("Claude 返回内容为空")
        candidate = self._extract_json_candidate(text)
        try:
            parsed = json.loads(candidate)
        except json.JSONDecodeError as exc:
            if repair_json is None:
                preview = text[:500].replace("\n", " ")
                raise RuntimeError(
                    "Claude returned invalid fallback JSON and json-repair is not installed; "
                    f"original parse error: {exc}; response preview: {preview}"
                ) from exc
            logger.warning("SDK文本兜底不是合法JSON，尝试json_repair: %s", exc)
            try:
                parsed = repair_json(candidate, return_objects=True)
            except Exception as repair_exc:
                preview = text[:500].replace("\n", " ")
                raise RuntimeError(
                    f"Claude返回内容无法修复为JSON: {repair_exc}; 原始内容前500字符: {preview}"
                ) from repair_exc
        parsed = self._apply_coco_business_corrections(parsed)
        validated = CocoJsonResponse.model_validate(parsed)
        return validated.model_dump_json(indent=2)

    def _apply_coco_business_corrections(self, parsed: dict[str, Any]) -> dict[str, Any]:
        """Apply deterministic rules that should not drift with LLM wording."""
        tickets = parsed.get("tickets")
        if not isinstance(tickets, list):
            return parsed

        for ticket in tickets:
            if not isinstance(ticket, dict):
                continue

            operation_party = str(ticket.get("operation_party") or "")
            assignee = str(ticket.get("assignee") or "")
            ticket_center = str(ticket.get("ticket_center") or "")
            primary_category = str(ticket.get("primary_category") or "")

            needs_center_dispatch = (
                "对应中心环境组" in operation_party
                or primary_category in {"堡垒机", "DFTM"}
            )
            center_unknown = ticket_center in {"", "无法判断", "未知", "-"}
            if needs_center_dispatch and center_unknown:
                missing_fields = ticket.get("missing_fields")
                if not isinstance(missing_fields, list):
                    missing_fields = []
                if "所属中心" not in missing_fields:
                    missing_fields.append("所属中心")

                ticket["ticket_center"] = "无法判断"
                ticket["missing_fields"] = missing_fields
                ticket["information_complete"] = False
                ticket["operation_party"] = "退回申请人"
                ticket["assignee"] = "退回申请人"
                ticket["review_result"] = "退回申请人"
                ticket["reason"] = "对应中心环境组工单缺少可判断的所属中心，无法按上海/苏州人员规则指派，需退回申请人补充。"
                continue

            if operation_party == "人工处理" or assignee == "人工处理":
                ticket["operation_party"] = "人工处理"
                ticket["assignee"] = "人工处理"
                ticket["review_result"] = "人工处理"

        return parsed

    def _extract_json_candidate(self, text: str) -> str:
        stripped = text.strip()
        if stripped.startswith("```"):
            lines = stripped.splitlines()
            if lines and lines[0].startswith("```"):
                lines = lines[1:]
            if lines and lines[-1].startswith("```"):
                lines = lines[:-1]
            stripped = "\n".join(lines).strip()

        object_start = stripped.find("{")
        object_end = stripped.rfind("}")
        if object_start != -1 and object_end != -1 and object_end > object_start:
            return stripped[object_start : object_end + 1]

        return stripped
