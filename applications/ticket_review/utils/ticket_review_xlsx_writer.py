from __future__ import annotations

import os
import tempfile
from copy import copy
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.utils.cell import coordinate_from_string


REVIEW_HEADERS = [
    ("行号", "row_number"),
    ("工单号", "ticket_no"),
    ("工单ID", "ticket_id"),
    ("实例ID", "instance_id"),
    ("类型", "ticket_type"),
    ("中心", "ticket_center"),
    ("标题", "title"),
    ("work_type", "work_type"),
    ("应用系统", "app_system"),
    ("一级分类", "primary_category"),
    ("二级分类", "secondary_category"),
    ("命中关键字", "matched_keyword"),
    ("初筛条件", "screening_condition"),
    ("必填字段", "required_fields"),
    ("缺失字段", "missing_fields"),
    ("信息完整", "information_complete"),
    ("操作方", "operation_party"),
    ("指派人员", "assignee"),
    ("结论", "review_result"),
    ("原因", "reason"),
]

APPEND_REVIEW_HEADERS = [
    ("预审-所属中心", "ticket_center"),
    ("预审-一级分类", "primary_category"),
    ("预审-二级分类", "secondary_category"),
    ("预审-命中关键词", "matched_keyword"),
    ("预审-初筛条件", "screening_condition"),
    ("预审-必填字段", "required_fields"),
    ("预审-缺失字段", "missing_fields"),
    ("预审-操作方", "operation_party"),
    ("预审-指派人员", "assignee"),
    ("预审-结论", "review_result"),
    ("预审-原因", "reason"),
]


def _cell_value(ticket: dict[str, Any], key: str) -> str:
    value = ticket.get(key, "")
    if isinstance(value, list):
        return "、".join(str(item) for item in value)
    if isinstance(value, bool):
        return "是" if value else "否"
    return str(value) if value is not None else ""


def ticket_review_to_xlsx(review_data: dict[str, Any], output_path: str | None = None) -> str:
    """将 coco 工单预审 JSON 结果导出为 XLSX。

    Args:
        review_data: CocoJsonResponse 解析后的字典。
        output_path: 输出路径，默认创建临时文件。

    Returns:
        生成的 xlsx 文件绝对路径。
    """
    if output_path is None:
        fd, output_path = tempfile.mkstemp(suffix=".xlsx", prefix="ticket_review_")
        os.close(fd)

    tickets = review_data.get("tickets", [])
    metadata = review_data.get("metadata", {})

    wb = Workbook()
    ws = wb.active
    ws.title = "预审结果"

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    cell_font = Font(name="微软雅黑", size=10)
    wrap_align = Alignment(wrap_text=True, vertical="top")
    center_align = Alignment(wrap_text=True, vertical="top", horizontal="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    pass_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    reject_fill = PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid")
    manual_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    widths = [8, 18, 18, 18, 12, 10, 35, 14, 16, 12, 12, 14, 20, 20, 20, 10, 14, 12, 10, 40]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    # 元信息
    ws.cell(row=1, column=1, value="预审结果汇总")
    ws.cell(row=1, column=1).font = Font(name="微软雅黑", size=14, bold=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(REVIEW_HEADERS))
    ws.cell(row=2, column=1, value=f"来源文件：{metadata.get('source_file', '')}")
    ws.cell(row=3, column=1, value=f"解析行数：{metadata.get('row_count', 0)}")
    ws.cell(row=4, column=1, value=f"使用技能：{metadata.get('skill_name', '')}")

    # 表头
    header_row = 6
    for col_idx, (header, _) in enumerate(REVIEW_HEADERS, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # 数据行
    for row_idx, ticket in enumerate(tickets, start=header_row + 1):
        review_result = ticket.get("review_result", "")
        if review_result == "通过":
            row_fill = pass_fill
        elif review_result == "退回申请人":
            row_fill = reject_fill
        else:
            row_fill = manual_fill

        for col_idx, (_, key) in enumerate(REVIEW_HEADERS, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=_cell_value(ticket, key))
            cell.font = cell_font
            cell.alignment = wrap_align if key in ("title", "screening_condition", "reason") else center_align
            cell.border = thin_border
            cell.fill = row_fill

        ws.row_dimensions[row_idx].height = max(24, _estimate_height(ticket))

    ws.freeze_panes = f"A{header_row + 1}"
    wb.save(output_path)
    return os.path.abspath(output_path)


def append_review_to_source_xlsx(
    source_bytes: bytes,
    review_data: dict[str, Any],
    output_path: str,
    keep_vba: bool = False,
) -> str:
    """Put review columns on the left while preserving the complete source workbook."""
    workbook = load_workbook(BytesIO(source_bytes), keep_vba=keep_vba)
    tickets = review_data.get("tickets", [])
    tickets_by_location = {
        (str(ticket.get("source_sheet") or ""), int(ticket.get("row_number") or 0)): ticket
        for ticket in tickets
    }

    for worksheet in workbook.worksheets:
        sheet_tickets = [
            ticket for ticket in tickets if str(ticket.get("source_sheet") or "") == worksheet.title
        ]
        if not sheet_tickets:
            continue

        header_row = min((int(ticket.get("row_number") or 2) for ticket in sheet_tickets), default=2) - 1
        header_row = max(header_row, 1)
        review_column_count = len(APPEND_REVIEW_HEADERS)

        # openpyxl moves cells but does not update all worksheet metadata when
        # inserting columns. Capture and restore the source layout explicitly.
        source_dimensions = {
            column_index_from_string(column_letter): copy(dimension)
            for column_letter, dimension in worksheet.column_dimensions.items()
        }
        table_ranges = {
            name: table if isinstance(table, str) else table.ref
            for name, table in worksheet.tables.items()
        }
        auto_filter_ref = worksheet.auto_filter.ref
        merged_ranges = [str(cell_range) for cell_range in worksheet.merged_cells.ranges]
        freeze_panes = worksheet.freeze_panes

        for merged_range in merged_ranges:
            worksheet.unmerge_cells(merged_range)
        worksheet.insert_cols(1, amount=review_column_count)

        for source_column, dimension in source_dimensions.items():
            target_letter = get_column_letter(source_column + review_column_count)
            target = worksheet.column_dimensions[target_letter]
            target.width = dimension.width
            target.hidden = dimension.hidden
            target.bestFit = dimension.bestFit
            target.outlineLevel = dimension.outlineLevel
            target.collapsed = dimension.collapsed

        for name, table_range in table_ranges.items():
            worksheet.tables[name].ref = _shift_range(table_range, review_column_count)
        if auto_filter_ref:
            worksheet.auto_filter.ref = _shift_range(auto_filter_ref, review_column_count)
        for merged_range in merged_ranges:
            worksheet.merge_cells(_shift_range(merged_range, review_column_count))

        if freeze_panes:
            coordinate = freeze_panes.coordinate if hasattr(freeze_panes, "coordinate") else str(freeze_panes)
            column_letter, row_number = coordinate_from_string(coordinate)
            shifted_column = column_index_from_string(column_letter) + review_column_count
            worksheet.freeze_panes = f"{get_column_letter(shifted_column)}{row_number}"
        else:
            worksheet.freeze_panes = f"{get_column_letter(review_column_count + 1)}{header_row + 1}"

        start_column = 1
        for offset, (header, _) in enumerate(APPEND_REVIEW_HEADERS):
            cell = worksheet.cell(row=header_row, column=start_column + offset, value=header)
            cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
            cell.border = Border(
                left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"), bottom=Side(style="thin"),
            )
            worksheet.column_dimensions[get_column_letter(start_column + offset)].width = 18

        for row_number in range(header_row + 1, worksheet.max_row + 1):
            ticket = tickets_by_location.get((worksheet.title, row_number))
            if not ticket:
                continue
            for offset, (_, key) in enumerate(APPEND_REVIEW_HEADERS):
                cell = worksheet.cell(
                    row=row_number,
                    column=start_column + offset,
                    value=_cell_value(ticket, key),
                )
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    workbook.save(output_path)
    return os.path.abspath(output_path)


def _shift_range(cell_range: str, column_offset: int) -> str:
    min_column, min_row, max_column, max_row = range_boundaries(cell_range)
    return (
        f"{get_column_letter(min_column + column_offset)}{min_row}:"
        f"{get_column_letter(max_column + column_offset)}{max_row}"
    )


def _estimate_height(ticket: dict[str, Any]) -> int:
    """根据可能换行的字段估算行高。"""
    max_lines = 1
    for key in ("title", "screening_condition", "reason"):
        value = ticket.get(key, "")
        if isinstance(value, str):
            lines = value.count("\n") + value.count("；") + value.count("。") + 1
            max_lines = max(max_lines, min(lines, 10))
    return max(24, max_lines * 16)
